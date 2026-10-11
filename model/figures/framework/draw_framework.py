"""Training framework figure for slides: which dataset supplies which inputs, and which inputs train which model.
One scene list, two back ends: SVG (rendered to PNG/PDF with Chrome) and PPTX with native, editable shapes grouped by panel.
usage: ~/miniconda3/bin/python draw_framework.py   (needs Pillow and python-pptx; fonts in ./fonts)
"""
import base64, io, json, math, os
from PIL import Image, ImageFont, ImageOps

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
A = json.load(open(os.path.join(HERE, "assets.json")))
W, H = 1920, 1080
FONT = "Inter"
FONT_FILES = {400: "Inter-Regular.ttf", 500: "Inter-Medium.ttf", 600: "Inter-SemiBold.ttf", 700: "Inter-Bold.ttf"}

# ---------------------------------------------------------------- palette: colour = data source, navy = learned model, coral = prediction
INK, BODY, MUTED, HAIR, PILL = "#18212F", "#475467", "#8B95A7", "#E2E6EC", "#F3F5F8"
NAVY, NAVY2, NAVY3, NAVY_L = "#1E2C52", "#3A4C82", "#6A7DB1", "#E7EBF5"
CORAL, CORAL_L = "#EB6547", "#FCE8E2"
DS = dict(cgm=("#0E9C88", "#E2F4F1", "#0A7466"), met=("#6E58D3", "#ECE9FB", "#4F3BB0"),
          big=("#DE9614", "#FBF1DC", "#9E6604"), col=("#CF5079", "#FAE6ED", "#A2355A"))

S = []                      # scene: list of dicts, drawn in order
GROUP = ["base"]


def g(name):
    GROUP[0] = name


def add(**kw):
    kw["group"] = GROUP[0]
    S.append(kw)


_fonts = {}


def tw(s, size, weight=400, ls=0.0):
    """text width in px (Inter metrics)"""
    key = (size, weight)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(os.path.join(HERE, "fonts", FONT_FILES[weight]), size)
    return _fonts[key].getlength(s) + ls * max(len(s) - 1, 0)


def rect(x, y, w, h, fill=None, stroke=None, sw=1.0, r=0, dash=None):
    add(t="rect", x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, r=r, dash=dash)


def poly(pts, fill=None, stroke=None, sw=1.0, closed=True, dash=None, cap="round"):
    add(t="poly", pts=[(round(a, 2), round(b, 2)) for a, b in pts], fill=fill, stroke=stroke, sw=sw, closed=closed, dash=dash, cap=cap)


def circle(cx, cy, r, fill=None, stroke=None, sw=1.0):
    add(t="circle", cx=cx, cy=cy, r=r, fill=fill, stroke=stroke, sw=sw)


def text(x, y, s, size=16, weight=400, fill=INK, anchor="start", ls=0.0):
    add(t="text", x=x, y=y, s=s, size=size, weight=weight, fill=fill, anchor=anchor, ls=ls, w=tw(s, size, weight, ls))


def image(x, y, w, h, path, r=0):
    add(t="image", x=x, y=y, w=w, h=h, path=path, r=r)


def line(pts, stroke, sw=2.0, dash=None, halo=False):
    if halo:                                              # white gap where this line passes over another
        poly(pts, stroke="#FFFFFF", sw=sw + 7, closed=False)
    poly(pts, stroke=stroke, sw=sw, closed=False, dash=dash)


def arrowhead(x, y, ang, color, size=9):
    """filled triangle with its tip at (x, y), pointing along ang (radians)"""
    c, s = math.cos(ang), math.sin(ang)
    b = (x - size * c, y - size * s)
    poly([(x, y), (b[0] - 0.55 * size * s, b[1] + 0.55 * size * c), (b[0] + 0.55 * size * s, b[1] - 0.55 * size * c)], fill=color)


def arrow(pts, color, sw=2.0, size=9, dash=None, halo=False):
    """polyline ending in an arrowhead at the last point"""
    (x0, y0), (x1, y1) = pts[-2], pts[-1]
    ang = math.atan2(y1 - y0, x1 - x0)
    trimmed = pts[:-1] + [(x1 - 0.8 * size * math.cos(ang), y1 - 0.8 * size * math.sin(ang))]
    line(trimmed, color, sw, dash, halo)
    arrowhead(x1, y1, ang, color, size)


def bez(p0, p1, p2, p3, n=28):
    return [((1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * p1[0] + 3 * (1 - t) * t ** 2 * p2[0] + t ** 3 * p3[0],
             (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * p1[1] + 3 * (1 - t) * t ** 2 * p2[1] + t ** 3 * p3[1])
            for t in [i / n for i in range(n + 1)]]


def ortho(pts, rad=12, n=7):
    """orthogonal polyline with rounded corners (subway style)"""
    out = [pts[0]]
    for i in range(1, len(pts) - 1):
        (ax, ay), (bx, by), (cx, cy) = pts[i - 1], pts[i], pts[i + 1]
        d1 = math.hypot(bx - ax, by - ay); d2 = math.hypot(cx - bx, cy - by)
        r = min(rad, d1 / 2, d2 / 2)
        p = (bx - (bx - ax) / d1 * r, by - (by - ay) / d1 * r); q = (bx + (cx - bx) / d2 * r, by + (cy - by) / d2 * r)
        out += [((1 - t) ** 2 * p[0] + 2 * (1 - t) * t * bx + t ** 2 * q[0], (1 - t) ** 2 * p[1] + 2 * (1 - t) * t * by + t ** 2 * q[1])
                for t in [k / n for k in range(n + 1)]]
    return out + [pts[-1]]


def chip(x, y, s, fg, bg, size=14, weight=500, h=26, pad=11):
    w = tw(s, size, weight) + 2 * pad
    rect(x, y, w, h, fill=bg, r=h / 2)
    text(x + pad, y + h / 2 + size * 0.36, s, size, weight, fg)
    return w


def plate(x, yc, w, h, dx, dy, front, side, top):
    """oblique box: thin front face, deep side face (a network layer)"""
    y0, y1 = yc - h / 2, yc + h / 2
    poly([(x + w, y0), (x + w + dx, y0 + dy), (x + w + dx, y1 + dy), (x + w, y1)], fill=side)
    poly([(x, y0), (x + dx, y0 + dy), (x + w + dx, y0 + dy), (x + w, y0)], fill=top)
    rect(x, y0, w, h, fill=front)


def spark(x, y, w, h, vals, color, sw=1.6, lo=None, hi=None):
    """mini curve scaled into the box"""
    v = [a for a in vals if a == a]
    mn, mx = (lo, hi) if lo is not None else (min(v), max(v))
    pts = [(x + w * i / (len(vals) - 1), y + h - h * (a - mn) / (mx - mn + 1e-9)) for i, a in enumerate(vals) if a == a]
    poly(pts, stroke=color, sw=sw, closed=False)
    return pts


def card(x, y, w, h, fill="#FFFFFF"):
    rect(x, y, w, h, fill=fill, stroke=HAIR, sw=1.5, r=16)


def station(x, y, color):
    circle(x, y, 6.5, fill=color, stroke="#FFFFFF", sw=2.5)


# ---------------------------------------------------------------- photo assets (square crops, so both back ends show the same)
def square(src, name, size=320):
    out = os.path.join(HERE, "assets", name)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    if not os.path.exists(out):
        ImageOps.fit(ImageOps.exif_transpose(Image.open(os.path.join(ROOT, src))).convert("RGB"), (size, size)).save(out, quality=90)
    return out


PHOTO = square(A["photo049"], "meal_photo_049.jpg")        # the example meal (photo branch and prediction)
PHOTO2 = square(A["photo"], "meal_photo_008.jpg")         # CGMacros card

# ================================================================ layout
X1, X1W = 48, 360                      # data column
RAIL = dict(big=428, met=444, col=460, cgm=476)      # dataset rails (left of the input list)
XP, XPW = 500, 400                     # input list
RA, RB = 924, 942                      # model rails (A / B; profile lane: logistic / ridge)
X3, X3W = 972, 480                     # models
X4, X4W = 1572, 300                    # prediction
ROW0, STEP = 338, 36                   # main input rows
PROW0 = 870                            # profile input rows

# ---------------------------------------------------------------- column headers
g("headers")
for x, w, num, lab in [(X1, X1W, "01", "DATA"), (XP - 72, XPW + 72 + 48, "02", "INPUTS"), (X3, X3W, "03", "MODELS"), (X4, X4W, "04", "PREDICTION")]:
    text(x, 50, num, 14, 700, CORAL, ls=1.5)
    text(x + 30, 50, lab, 14, 600, MUTED, ls=2.5)
    rect(x, 64, w, 1.5, fill=HAIR)
    rect(x, 63.25, 22, 3, fill=CORAL)

# ---------------------------------------------------------------- datasets
g("data")


def ds_card(y, h, key, name, sub, chips, official=False, extra=None):
    col, light, dark = DS[key]
    card(X1, y, X1W, h)
    rect(X1, y + 18, 5, 34, fill=col, r=2.5)
    text(X1 + 24, y + 42, name, 24, 700, INK)
    if official:
        bx = X1 + 24 + tw(name, 24, 700) + 10
        chip(bx, y + 22, "official", MUTED, PILL, 12, 600, 22, 9)
    text(X1 + 24, y + 66, sub, 15, 400, BODY)
    cx, cy = X1 + 24, y + (extra or 82)
    for c in chips:
        cw = tw(c, 14, 500) + 22
        if cx + cw > X1 + X1W - 20:
            cx, cy = X1 + 24, cy + 34
        chip(cx, cy, c, dark, light)
        cx += cw + 8
    return cy + 26


# CGMacros: photo + one day of CGM + chips + train/test timeline
yc0 = 92
card(X1, yc0, X1W, 300)
rect(X1, yc0 + 18, 5, 34, fill=DS["cgm"][0], r=2.5)
text(X1 + 24, yc0 + 42, "CGMacros", 24, 700, INK)
chip(X1 + 24 + tw("CGMacros", 24, 700) + 10, yc0 + 22, "official", MUTED, PILL, 12, 600, 22, 9)
text(X1 + 24, yc0 + 66, "45 adults · 1,397 meals · ~10 days", 15, 400, BODY)
image(X1 + 24, yc0 + 84, 84, 84, PHOTO2, r=10)
gx, gy, gw, gh = X1 + 124, yc0 + 88, X1W - 148, 76            # one day of CGM with meal marks
for k in range(4):
    rect(gx, gy + gh * k / 3, gw, 1, fill="#EEF1F5")
day = A["hist"][2] + A["hist"][5] + A["hist"][7]
spark(gx, gy + 4, gw, gh - 8, [v for v in day], DS["cgm"][0], 1.8)
for k in (0, 25, 50):
    circle(gx + gw * k / (len(day) - 1), gy + gh + 1, 3.2, fill=DS["cgm"][0])
cy = yc0 + 184
cx = X1 + 24
for c in ["Meal photos", "CGM", "Heart rate", "Activity", "Food log", "Age · sex · BMI"]:
    cw = tw(c, 14, 500) + 22
    if cx + cw > X1 + X1W - 20:
        cx, cy = X1 + 24, cy + 34
    chip(cx, cy, c, DS["cgm"][2], DS["cgm"][1]); cx += cw + 8
ty = yc0 + 262                                               # timeline: days 1-7 train, 8+ test
for d in range(10):
    rect(X1 + 24 + d * 31.2, ty, 28, 8, fill=DS["cgm"][0] if d < 7 else "#FFFFFF", stroke=None if d < 7 else DS["cgm"][0], sw=1.5, r=4)
text(X1 + 24, ty + 26, "days 1–7 · train", 13, 500, DS["cgm"][2])
text(X1 + 24 + 7 * 31.2, ty + 26, "8+ · test", 13, 500, MUTED)

ds_card(410, 168, "met", "MetaboNet", "1,085 people with type 1 diabetes · 671k meals", ["CGM", "Carbs", "Insulin", "Age · sex · BMI"], True)
ds_card(596, 140, "big", "BIG IDEAs", "16 adults · 397 meals", ["CGM", "Food log", "Sex"])
ds_card(842, 174, "col", "Colas 2019", "208 adults · 2.7-year follow-up", ["CGM, 1–2 days", "Clinical labs", "T2D outcome"])

# ---------------------------------------------------------------- inputs: photo branch
g("photo")
PY = 188                                                     # photo branch centre line
text(XP, 112, "Meal photo", 15, 600, INK)
text(XP, 129, "CGMacros", 13, 400, MUTED)
image(XP, PY - 48, 96, 96, PHOTO, r=12)
arrow([(XP + 104, PY), (XP + 126, PY)], MUTED, 1.8, 8)
ex = XP + 134
for i, (hh, dd) in enumerate([(96, 46), (80, 38), (64, 30), (48, 22)]):
    plate(ex + i * 27, PY + 6, 9, hh, dd * 0.62, -dd * 0.42, "#7E8DAA", "#A8B4CB", "#C9D1E0")
text(ex + 62, 112, "Image encoder", 15, 600, INK, "middle")
text(ex + 62, 129, "MobileCLIP-S2 · frozen", 13, 400, MUTED, "middle")
arrow([(ex + 130, PY), (ex + 150, PY)], MUTED, 1.8, 8)
vx = ex + 158
for k in range(8):                                           # embedding vector
    shade = ["#D7DDEA", "#9AA8C6", "#C2CBDD", "#6F82AB", "#E2E7F0", "#8C9BBC", "#B3BED5", "#7D8FB5"][k]
    rect(vx, PY - 48 + k * 12, 14, 12, fill=shade)
rect(vx, PY - 48, 14, 96, stroke="#6F82AB", sw=1.2)
text(vx + 7, PY + 66, "512-d", 13, 500, MUTED, "middle")
arrow([(vx + 22, PY), (vx + 42, PY)], MUTED, 1.8, 8)
gx0 = vx + 52                                                # nutrition estimator: nearest meals in a library
for i in range(3):
    for j in range(3):
        hit = (i, j) in [(0, 1), (2, 2)]
        rect(gx0 + j * 22, PY - 31 + i * 22, 17, 17, fill=DS["cgm"][0] if hit else "#E4E8EF", r=4)
text(gx0 + 28, 112, "Nutrition estimator", 15, 600, INK, "middle")
text(gx0 + 28, 129, "nearest meals + GBM", 13, 400, MUTED, "middle")
arrow([(gx0 + 28, PY + 40), (gx0 + 28, ROW0 - 16)], MUTED, 1.8, 8)
text(gx0 + 18, PY + 98, "carbs · protein · fat · fiber · kcal", 12, 500, MUTED, "end")

# ---------------------------------------------------------------- inputs: rows (which dataset supplies them, which model uses them)
ROWS = [("Nutrients from photo", "cgm", "AB", False),
        ("Glucose before the meal", "cgm met big", "AB", False),
        ("Meal time · previous meal", "cgm met big", "AB", False),
        ("First-week glucose profile", "cgm big", "A", False),
        ("Similar meals, own and others'", "cgm", "A", False),
        ("Past meals with their glucose curves", "cgm met", "B", False),
        ("Heart rate", "cgm", "AB", False),
        ("Activity (METs, active kcal)", "cgm", "A", False),
        ("Age · sex · BMI", "cgm met big", "AB", False),
        ("Logged kcal (optional)", "cgm big", "A", False),
        ("Insulin dose and timing", "met", "B", False),
        ("Glucose after the meal (target)", "cgm met big", "AB", True)]
PROWS = [("CGM metrics (up to 7 days)", "cgm big col", "LR", False),
         ("Age · BMI · fasting glucose · HbA1c", "col", "L", False),
         ("Lab HbA1c (target)", "cgm big col", "R", True),
         ("Developed type 2 diabetes (target)", "col", "L", True)]


def rows_y(rows, y0):
    return [y0 + i * STEP for i in range(len(rows))]


MY, PYS = rows_y(ROWS, ROW0), rows_y(PROWS, PROW0)
g("inputs")
for (lab, src, use, target), y in list(zip(ROWS, MY)) + list(zip(PROWS, PYS)):
    if target:
        rect(XP, y - 14, XPW, 28, fill="#FFFFFF", stroke="#C3CAD5", sw=1.3, r=14, dash="5,4")
    else:
        rect(XP, y - 14, XPW, 28, fill=PILL, r=14)
    text(XP + 16, y + 5.5, lab, 15, 500, BODY if target else INK)

# lane label for the profile models
g("lanes")
rect(X1, 812, W - 2 * X1, 1.2, fill=HAIR)
text(X1, 830, "PROFILE MODELS", 12, 600, MUTED, ls=2.0)

# ---------------------------------------------------------------- dataset rails (subway map) with connectors from the cards
g("rails_data")
allrows = list(zip(ROWS, MY)) + list(zip(PROWS, PYS))
ports = dict(cgm=ROW0, met=500, big=678, col=PROW0 + STEP)        # where each card's connector joins its rail
for key in ["big", "met", "col", "cgm"]:
    ys = [y for (lab, src, use, t), y in allrows if key in src.split()]
    x = RAIL[key]; col = DS[key][0]
    top, bot = min(ys + [ports[key]]), max(ys + [ports[key]])
    line([(x, top), (x, bot)], col, 3.2)
for key in ["big", "met", "col", "cgm"]:                     # connectors over the rails (white gap where they cross)
    x = RAIL[key]; col = DS[key][0]
    line(ortho([(X1 + X1W, ports[key]), (x, ports[key])]), col, 3.2, halo=True)
line([(X1 + X1W, PY), (XP - 8, PY)], DS["cgm"][0], 3.2)      # CGMacros photos -> photo branch
arrowhead(XP - 2, PY, 0, DS["cgm"][0], 10)
for key in ["big", "met", "col", "cgm"]:
    for (lab, src, use, t), y in allrows:
        if key in src.split():
            station(RAIL[key], y, DS[key][0])

# ---------------------------------------------------------------- model rails
g("rails_model")
A_PORT, B_PORT, R_PORT, L_PORT = 232, 590, 874, 950
for xr, letter, rows, port in [(RA, "A", list(zip(ROWS, MY)), A_PORT), (RB, "B", list(zip(ROWS, MY)), B_PORT),
                               (RA, "L", list(zip(PROWS, PYS)), L_PORT), (RB, "R", list(zip(PROWS, PYS)), R_PORT)]:
    ys = [y for (lab, src, use, t), y in rows if letter in use]
    top, bot = min(ys + [port]), max(ys + [port])
    line([(xr, top), (xr, bot)], NAVY, 3.0)
    arrow(ortho([(xr, port), (X3 - 2, port)]), NAVY, 3.0, 11)
    for y in ys:
        circle(xr, y, 6.5, fill=NAVY, stroke="#FFFFFF", sw=2.5)
for xr, letter in [(RA, "A"), (RB, "B")]:
    text(xr, MY[-1] + 30, letter, 12, 700, NAVY, "middle")
for (lab, src, use, t), y in list(zip(ROWS, MY)) + list(zip(PROWS, PYS)):      # thin tick from each row to its first model rail
    xs = [RA if (c in "AL") else RB for c in use]
    line([(XP + XPW, y), (min(xs) - 7, y)], "#C3CAD5", 1.4)

# ---------------------------------------------------------------- model A: gradient-boosted trees
g("model_A")
card(X3, 92, X3W, 280)
text(X3 + 24, 124, "MODEL A", 12, 700, NAVY3, ls=2.0)
text(X3 + 24, 152, "Gradient-boosted trees", 22, 700, INK)
text(X3 + 24, 174, "LightGBM · glucose at 15–120 min, peak, iAUC", 14, 400, MUTED)
LEAF = ["#D4DBEA", "#8D9DC2", "#5A6FA0", "#2E4170", "#B7C3DC", "#6B80B0", "#A4B2D1", "#3D5285"]
for t_i, tx in enumerate([X3 + 78, X3 + 200, X3 + 322]):
    root, l1, l2 = (tx, 200), [(tx - 34, 230), (tx + 34, 230)], [(tx - 51, 262), (tx - 17, 262), (tx + 17, 262), (tx + 51, 262)]
    for p in l1:
        line([root, p], NAVY3, 1.6)
    for k, p in enumerate(l2):
        line([l1[k // 2], p], NAVY3, 1.6)
    circle(*root, 7, fill=NAVY)
    for p in l1:
        circle(*p, 6, fill=NAVY2)
    for k, p in enumerate(l2):
        rect(p[0] - 7, p[1] - 7, 14, 14, fill=LEAF[(k + 3 * t_i) % 8], r=3.5)
    if t_i < 2:
        text(tx + 61, 238, "+", 20, 500, MUTED, "middle")
text(X3 + 402, 238, "…", 22, 600, MUTED, "middle")
text(X3 + 24, 336, "Trained on", 13, 500, MUTED)
cx = X3 + 24 + tw("Trained on", 13, 500) + 10
cx += chip(cx, 318, "CGMacros · days 1–7", DS["cgm"][2], DS["cgm"][1], 13) + 8
chip(cx, 318, "BIG IDEAs", DS["big"][2], DS["big"][1], 13)

# ---------------------------------------------------------------- model B: in-context transformer
g("model_B")
YB = 392
card(X3, YB, X3W, 404)
text(X3 + 24, YB + 32, "MODEL B", 12, 700, NAVY3, ls=2.0)
text(X3 + 24, YB + 60, "In-context transformer", 22, 700, INK)
text(X3 + 24, YB + 82, "reads the person's past meals, outputs the whole curve", 14, 400, MUTED)
tx0, ty0 = X3 + 26, YB + 116                                 # meal tokens: new meal + past meals with their curves
text(tx0 + 25, ty0 - 10, "meal tokens", 12, 600, MUTED, "middle")
rect(tx0, ty0, 50, 28, fill=CORAL_L, stroke=CORAL, sw=1.6, r=7)
text(tx0 + 25, ty0 + 19, "new", 12, 600, CORAL, "middle")
for k in range(3):
    yy = ty0 + 38 + k * 36
    src = "met" if k == 2 else "cgm"
    rect(tx0, yy, 50, 28, fill="#FFFFFF", stroke=DS[src][0], sw=1.4, r=7)
    spark(tx0 + 7, yy + 5, 36, 17, A["hist"][[0, 3, 6, 4][k]], DS[src][0], 1.5)
for k in range(3):
    circle(tx0 + 25, ty0 + 152 + k * 7, 1.8, fill=MUTED)
text(tx0 + 25, ty0 + 188, "up to 30", 12, 500, MUTED, "middle")
arrow([(tx0 + 58, B_PORT), (tx0 + 84, B_PORT)], MUTED, 1.8, 8)
px0 = tx0 + 98
for k in range(3):                                           # transformer layers
    plate(px0 + k * 50, B_PORT + 8, 11, 140, 34, -20, NAVY, NAVY2, NAVY3)
text(px0 + 62, B_PORT + 104, "Transformer layers × 3", 13, 600, BODY, "middle")
arrow([(px0 + 150, B_PORT), (px0 + 168, B_PORT)], MUTED, 1.8, 8)
hx = px0 + 176
plate(hx, B_PORT + 6, 10, 84, 24, -14, CORAL, "#F2957F", "#F7B9A9")
text(hx + 17, B_PORT + 72, "quantile head", 13, 600, BODY, "middle")
arrow([(hx + 40, B_PORT), (hx + 58, B_PORT)], MUTED, 1.8, 8)
fx, fy, fw, fh = hx + 66, B_PORT - 44, X3 + X3W - 24 - (hx + 66), 86     # output: curve with 10-90% band
lo, hi, md = A["lo"], A["hi"], A["pred"]
mn, mx = min(lo), max(hi)
band = [(fx + fw * i / 24, fy + fh - fh * (v - mn) / (mx - mn)) for i, v in enumerate(hi)] + \
       [(fx + fw * i / 24, fy + fh - fh * (v - mn) / (mx - mn)) for i, v in reversed(list(enumerate(lo)))]
poly(band, fill=NAVY_L)
spark(fx, fy, fw, fh, md, NAVY, 2.4, mn, mx)
rect(fx, fy + fh + 4, fw, 1.2, fill="#C9D0DC")
text(fx + fw / 2, B_PORT + 72, "10 / 50 / 90%", 13, 600, BODY, "middle")
for k, (lab, key) in enumerate([("Pretrain · MetaboNet, 671k meals", "met"), ("Fine-tune · CGMacros, days 1–7", "cgm")]):
    yy = YB + 330 + k * 34
    circle(X3 + 34, yy + 13, 10, fill=DS[key][0])
    text(X3 + 34, yy + 17.5, str(k + 1), 12, 700, "#FFFFFF", "middle")
    chip(X3 + 52, yy, lab, DS[key][2], DS[key][1], 13)

# ---------------------------------------------------------------- profile models
g("profile_models")
for y0, title, sub, port, kind in [(842, "HbA1c estimate", "Ridge regression · 261 people", R_PORT, "ridge"),
                                   (918, "Type 2 diabetes risk", "Logistic regression · 208 people", L_PORT, "logit")]:
    card(X3, y0, X3W, 64)
    text(X3 + 24, y0 + 28, title, 17, 700, INK)
    text(X3 + 24, y0 + 48, sub, 13, 400, MUTED)
    gx1, gy1 = X3 + X3W - 112, y0 + 12
    if kind == "ridge":
        for (a, b) in [(6, 32), (16, 28), (24, 30), (34, 20), (44, 22), (52, 14), (64, 12), (72, 6), (84, 8)]:
            circle(gx1 + a, gy1 + b, 3, fill=NAVY3)
        line([(gx1, gy1 + 36), (gx1 + 90, gy1 + 2)], NAVY, 2.2)
    else:
        pts = [(gx1 + 90 * i / 30, gy1 + 38 - 36 / (1 + math.exp(-(i - 15) / 3))) for i in range(31)]
        line(pts, NAVY, 2.4)
        rect(gx1, gy1 + 39, 90, 1.2, fill="#C9D0DC")

# ---------------------------------------------------------------- blend and prediction
g("prediction")
BX, BY = 1512, 430
arrow(bez((X3 + X3W, A_PORT), (X3 + X3W + 40, A_PORT), (BX - 30, BY - 70), (BX - 17, BY - 17)), NAVY, 2.6, 10)
arrow(bez((X3 + X3W, B_PORT), (X3 + X3W + 40, B_PORT), (BX - 30, BY + 70), (BX - 17, BY + 17)), NAVY, 2.6, 10)
circle(BX, BY, 24, fill="#FFFFFF", stroke=NAVY, sw=2.4)
text(BX, BY + 7.5, "Σ", 21, 600, NAVY, "middle")
text(BX + 8, BY + 46, "blend", 13, 600, MUTED)
arrow([(BX + 26, BY), (X4 - 2, BY)], NAVY, 2.6, 10)
card(X4, 92, X4W, 704)                                       # prediction for the example meal
rect(X4, 110, 5, 34, fill=CORAL, r=2.5)
text(X4 + 24, 134, "Glucose response", 21, 700, INK)
text(X4 + 24, 156, "next 2 hours · 80% band", 13, 400, MUTED)
cx0, cy0, cw0, ch0 = X4 + 54, 198, X4W - 74, 272               # chart: 60 min of real CGM before the meal, then the prediction
pre = A["pre"][-13:]
lo_v, hi_v = 80, 270
yv = lambda v: cy0 + ch0 - ch0 * (v - lo_v) / (hi_v - lo_v)
xv = lambda m: cx0 + cw0 * (m + 60) / 180
for v in (100, 150, 200, 250):
    rect(cx0, yv(v), cw0, 1, fill="#EEF1F5")
    text(cx0 - 8, yv(v) + 4, str(v), 11, 500, MUTED, "end")
text(cx0 - 8, cy0 - 12, "mg/dL", 11, 500, MUTED, "start")
band = [(xv(5 * i), yv(v)) for i, v in enumerate(hi)] + [(xv(5 * i), yv(v)) for i, v in reversed(list(enumerate(lo)))]
poly(band, fill=CORAL_L)
poly([(xv(-60 + 5 * i), yv(v)) for i, v in enumerate(pre)], stroke="#9AA3B2", sw=2.4, closed=False)
line([(xv(0), cy0 + 8), (xv(0), cy0 + ch0)], "#B8C0CC", 1.4, dash="4,4")
poly([(xv(5 * i), yv(v)) for i, v in enumerate(md)], stroke=CORAL, sw=3.2, closed=False)
image(xv(0) - 17, cy0 - 26, 34, 34, PHOTO, r=17)
for m, lab in [(-60, "−60"), (0, "0"), (60, "60"), (120, "120 min")]:
    text(xv(m), cy0 + ch0 + 20, lab, 12, 500, MUTED, "middle" if m < 120 else "end")
text(xv(-30), yv(max(pre)) - 12, "CGM", 11, 600, "#9AA3B2", "middle")
B = A["blend"]
tiles = [("Peak glucose", f"{B['Gmax']:.0f}", "mg/dL", "real 187"), ("Peak rise", f"+{B['dpeak']:.0f}", "mg/dL", "real +92"),
         ("Time to peak", f"{B['ttp']:.0f}", "min", "real 106"), ("iAUC, mg/dL·min", f"{B['iauc'] / 1000:.1f}k", "", "real 6.3k")]
for k, (lab, val, unit, real) in enumerate(tiles):
    tx_, ty_ = X4 + 24 + (k % 2) * 132, 528 + (k // 2) * 92
    rect(tx_, ty_, 120, 80, fill="#F7F8FA", r=12)
    text(tx_ + 14, ty_ + 24, lab, 12, 500, MUTED)
    text(tx_ + 14, ty_ + 52, val, 22, 700, INK)
    text(tx_ + 18 + tw(val, 22, 700), ty_ + 52, unit, 11, 500, MUTED)
    text(tx_ + 14, ty_ + 70, real, 11, 500, MUTED)
text(X4 + 24, 744, "Example: a real lunch on an unseen", 12, 500, MUTED)
text(X4 + 24, 762, "test day; \"real\" is the recorded CGM.", 12, 500, MUTED)
for y, lab in [(R_PORT, "Estimated HbA1c"), (L_PORT, "Risk tier")]:
    arrow([(X3 + X3W, y), (X4 - 2, y)], NAVY, 2.4, 10)
    rect(X4, y - 18, X4W, 36, fill=CORAL_L, r=18)
    text(X4 + 20, y + 5.5, lab, 15, 600, "#B2432A")

# ================================================================ back end 1: SVG


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_svg(path):
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" viewBox="0 0 {W} {H}" width="{W}" height="{H}" '
           f'font-family="{FONT}, \'Helvetica Neue\', Arial, sans-serif">', f'<rect width="{W}" height="{H}" fill="#FFFFFF"/>']
    clip = 0
    cur = None
    for it in S:
        if it["group"] != cur:
            if cur is not None:
                out.append("</g>")
            out.append(f'<g id="{it["group"]}">'); cur = it["group"]
        st = lambda: (f' fill="{it["fill"] or "none"}"' + (f' stroke="{it["stroke"]}" stroke-width="{it["sw"]}"' if it.get("stroke") else "")
                      + (f' stroke-dasharray="{it["dash"]}"' if it.get("dash") else ""))
        if it["t"] == "rect":
            out.append(f'<rect x="{it["x"]:.2f}" y="{it["y"]:.2f}" width="{it["w"]:.2f}" height="{it["h"]:.2f}" rx="{it["r"]}"{st()}/>')
        elif it["t"] == "circle":
            out.append(f'<circle cx="{it["cx"]:.2f}" cy="{it["cy"]:.2f}" r="{it["r"]}"{st()}/>')
        elif it["t"] == "poly":
            d = "M" + " L".join(f"{a:.2f},{b:.2f}" for a, b in it["pts"]) + (" Z" if it["closed"] else "")
            extra = ' stroke-linejoin="round" stroke-linecap="round"' if it.get("stroke") else ""
            out.append(f'<path d="{d}"{st()}{extra}/>')
        elif it["t"] == "text":
            anc = {"start": "start", "middle": "middle", "end": "end"}[it["anchor"]]
            ls = f' letter-spacing="{it["ls"]}"' if it["ls"] else ""
            out.append(f'<text x="{it["x"]:.2f}" y="{it["y"]:.2f}" font-size="{it["size"]}" font-weight="{it["weight"]}" fill="{it["fill"]}" '
                       f'text-anchor="{anc}"{ls} xml:space="preserve">{esc(it["s"])}</text>')
        elif it["t"] == "image":
            b64 = base64.b64encode(open(it["path"], "rb").read()).decode()
            clip += 1
            out.append(f'<clipPath id="c{clip}"><rect x="{it["x"]}" y="{it["y"]}" width="{it["w"]}" height="{it["h"]}" rx="{it["r"]}"/></clipPath>')
            out.append(f'<image x="{it["x"]}" y="{it["y"]}" width="{it["w"]}" height="{it["h"]}" clip-path="url(#c{clip})" preserveAspectRatio="xMidYMid slice" '
                       f'href="data:image/jpeg;base64,{b64}"/>')
    out.append("</g></svg>")
    open(path, "w").write("\n".join(out))


# ================================================================ back end 2: PPTX with native shapes


def to_pptx(path):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
    from pptx.enum.dml import MSO_LINE
    from pptx.oxml.ns import qn
    E = 6350                                                     # EMU per px (1920 px = 13.333 in)
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(W * E), Emu(H * E)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))

    def clean(shape):
        st = shape._element.find(qn("p:style"))
        if st is not None:
            shape._element.remove(st)

    def style(shape, it):
        clean(shape)
        if it.get("fill"):
            shape.fill.solid(); shape.fill.fore_color.rgb = rgb(it["fill"])
        else:
            shape.fill.background()
        if it.get("stroke"):
            shape.line.color.rgb = rgb(it["stroke"]); shape.line.width = Emu(int(it["sw"] * E))
            if it.get("dash"):
                shape.line.dash_style = MSO_LINE.DASH
            shape.line._get_or_add_ln().set("cap", "rnd")
        else:
            shape.line.fill.background()

    groups = {}
    order = []
    for it in S:
        if it["group"] not in groups:
            groups[it["group"]] = slide.shapes.add_group_shape(); order.append(it["group"])
        sh = groups[it["group"]].shapes
        t = it["t"]
        if t == "rect":
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if it["r"] else MSO_SHAPE.RECTANGLE
            s = sh.add_shape(kind, Emu(int(it["x"] * E)), Emu(int(it["y"] * E)), Emu(max(int(it["w"] * E), 1)), Emu(max(int(it["h"] * E), 1)))
            if it["r"]:
                s.adjustments[0] = min(0.5, it["r"] / max(min(it["w"], it["h"]), 1e-6))
            style(s, it)
        elif t == "circle":
            r = it["r"]
            s = sh.add_shape(MSO_SHAPE.OVAL, Emu(int((it["cx"] - r) * E)), Emu(int((it["cy"] - r) * E)), Emu(int(2 * r * E)), Emu(int(2 * r * E)))
            style(s, it)
        elif t == "poly":
            pts = it["pts"]
            fb = sh.build_freeform(Emu(int(pts[0][0] * E)), Emu(int(pts[0][1] * E)), scale=1.0)
            fb.add_line_segments([(Emu(int(a * E)), Emu(int(b * E))) for a, b in pts[1:]], close=it["closed"])
            s = fb.convert_to_shape()
            style(s, it)
        elif t == "text":
            size = it["size"]; w = it["w"] + 6
            x = it["x"] - (w / 2 if it["anchor"] == "middle" else (w - 3 if it["anchor"] == "end" else 0))
            tb = sh.add_textbox(Emu(int(x * E)), Emu(int((it["y"] - 0.98 * size) * E)), Emu(int(w * E)), Emu(int(1.3 * size * E)))
            tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[it["anchor"]]
            r = p.add_run(); r.text = it["s"]
            f = r.font; f.name = FONT; f.size = Pt(size * 0.5); f.bold = it["weight"] >= 600; f.color.rgb = rgb(it["fill"])
            if it["ls"]:
                r._r.get_or_add_rPr().set("spc", str(int(it["ls"] * 0.5 * 100)))
        elif t == "image":
            s = sh.add_picture(it["path"], Emu(int(it["x"] * E)), Emu(int(it["y"] * E)), Emu(int(it["w"] * E)), Emu(int(it["h"] * E)))
            if it["r"]:
                geom = s._element.spPr.find(qn("a:prstGeom"))
                geom.set("prst", "roundRect")
    for name in order:
        groups[name].name = name
    prs.save(path)


if __name__ == "__main__":
    os.makedirs(os.path.join(HERE, "out"), exist_ok=True)
    to_svg(os.path.join(HERE, "out", "training_framework.svg"))
    to_pptx(os.path.join(HERE, "out", "training_framework.pptx"))
    print("scene items:", len(S))
